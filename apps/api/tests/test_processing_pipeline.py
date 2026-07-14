from __future__ import annotations

import io

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter

from app.services.files.mime_registry import Processor
from app.services.files.parsers import ExtractedSection
from app.services.files.pipeline import (
    FileProcessingPipeline,
    ProcessorFailure,
)


def jpeg_bytes() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (16, 12), "white").save(output, format="JPEG")
    return output.getvalue()


def blank_pdf() -> bytes:
    output = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(output)
    return output.getvalue()


class SuccessfulOCR:
    async def extract(self, content: bytes, mime_type: str) -> list[ExtractedSection]:
        assert mime_type in {"image/jpeg", "application/pdf"}
        return [ExtractedSection("Product | Origin | Price\nMango | India | 10")]


class UnavailableOCR:
    def __init__(self) -> None:
        self.calls = 0

    async def extract(self, content: bytes, mime_type: str) -> list[ExtractedSection]:
        self.calls += 1
        raise ProcessorFailure("provider_timeout", retryable=True)


async def test_market_jpeg_normalizes_then_uses_document_ai() -> None:
    pipeline = FileProcessingPipeline({Processor.document_ai: SuccessfulOCR()})
    result = await pipeline.process(
        jpeg_bytes(), "WhatsApp Market Report.JPEG", "application/octet-stream"
    )
    assert result.processor == Processor.document_ai
    assert result.derivative is not None
    assert result.derivative.mime_type == "image/jpeg"
    assert result.sections[0].text.startswith("Product")
    assert result.fallback_used is False


async def test_scanned_pdf_uses_ocr_when_native_parser_has_no_text() -> None:
    pipeline = FileProcessingPipeline({Processor.document_ai: SuccessfulOCR()})
    result = await pipeline.process(blank_pdf(), "scan.pdf", "application/pdf")
    assert result.processor == Processor.document_ai
    assert result.sections


async def test_retry_is_bounded_then_uses_compatible_fallback() -> None:
    unavailable = UnavailableOCR()
    pipeline = FileProcessingPipeline(
        {
            Processor.document_ai: unavailable,
            Processor.gemini_vision: SuccessfulOCR(),
        }
    )
    result = await pipeline.process(jpeg_bytes(), "market.jpeg", "image/jpeg")
    assert unavailable.calls == 3
    assert result.fallback_used is True
    assert result.processor == Processor.gemini_vision
    assert len(result.attempts) == 4


async def test_docx_extracts_paragraphs_and_tables_natively() -> None:
    document = Document()
    document.add_paragraph("Payment terms")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Advance"
    table.cell(1, 1).text = "30%"
    output = io.BytesIO()
    document.save(output)
    result = await FileProcessingPipeline().process(
        output.getvalue(), "terms.docx", "application/octet-stream"
    )
    assert any("Payment terms" in section.text for section in result.sections)
    assert any("30%" in section.text for section in result.sections)


async def test_xlsx_preserves_sheet_headers_and_row_context() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Market"
    sheet.append(["Product", "Price"])
    sheet.append(["Mango", 10])
    output = io.BytesIO()
    workbook.save(output)
    result = await FileProcessingPipeline().process(
        output.getvalue(), "prices.xlsx", "application/octet-stream"
    )
    assert result.sections[0].metadata == {"sheet": "Market", "row": 2}
    assert "Product: Mango" in result.sections[0].text
