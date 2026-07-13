from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field

import xlrd
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


@dataclass(frozen=True)
class ExtractedSection:
    text: str
    metadata: dict[str, str | int] = field(default_factory=dict)


def parse_text(content: bytes) -> list[ExtractedSection]:
    return [ExtractedSection(content.decode("utf-8-sig"))]


def parse_json(content: bytes) -> list[ExtractedSection]:
    parsed = json.loads(content.decode("utf-8-sig"))
    return [ExtractedSection(json.dumps(parsed, ensure_ascii=False, indent=2))]


def parse_csv(content: bytes) -> list[ExtractedSection]:
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    if not rows:
        return []
    header = rows[0]
    return [
        ExtractedSection(
            " | ".join(
                f"{header[index] if index < len(header) else index}: {value}"
                for index, value in enumerate(row)
            ),
            {"row": row_number},
        )
        for row_number, row in enumerate(rows[1:], start=2)
    ]


def parse_docx(content: bytes) -> list[ExtractedSection]:
    document = Document(io.BytesIO(content))
    sections = [
        ExtractedSection(p.text, {"kind": "paragraph"})
        for p in document.paragraphs
        if p.text.strip()
    ]
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            sections.append(
                ExtractedSection(
                    " | ".join(cell.text for cell in row.cells),
                    {"kind": "table", "table": table_index, "row": row_index},
                )
            )
    return sections


def parse_xlsx(content: bytes) -> list[ExtractedSection]:
    workbook = load_workbook(
        io.BytesIO(content), read_only=True, data_only=False, keep_links=False
    )
    sections: list[ExtractedSection] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(value or "") for value in rows[0]]
        for row_number, row in enumerate(rows[1:], start=2):
            values = [str(value) if value is not None else "" for value in row]
            sections.append(
                ExtractedSection(
                    " | ".join(
                        f"{headers[index] or index + 1}: {value}"
                        for index, value in enumerate(values)
                    ),
                    {"sheet": sheet.title, "row": row_number},
                )
            )
    return sections


def parse_xls(content: bytes) -> list[ExtractedSection]:
    workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    sections: list[ExtractedSection] = []
    for sheet in workbook.sheets():
        if sheet.nrows == 0:
            continue
        headers = [str(sheet.cell_value(0, column)) for column in range(sheet.ncols)]
        for row_number in range(1, sheet.nrows):
            values = [
                str(sheet.cell_value(row_number, column))
                for column in range(sheet.ncols)
            ]
            text = " | ".join(
                f"{headers[index] or index + 1}: {value}"
                for index, value in enumerate(values)
            )
            sections.append(
                ExtractedSection(text, {"sheet": sheet.name, "row": row_number + 1})
            )
    workbook.release_resources()
    return sections


def parse_pdf(content: bytes) -> list[ExtractedSection]:
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise ValueError("encrypted_file")
    return [
        ExtractedSection(text, {"page": number})
        for number, page in enumerate(reader.pages, start=1)
        if (text := (page.extract_text() or "").strip())
    ]


def parse_pptx(content: bytes) -> list[ExtractedSection]:
    presentation = Presentation(io.BytesIO(content))
    sections: list[ExtractedSection] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        )
        if text:
            sections.append(ExtractedSection(text, {"slide": slide_number}))
    return sections
